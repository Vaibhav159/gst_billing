"""Cancelled bills leave every figure (plan 1A, Task 2): one test per place that adds bills up
(server research S§2, M1-M18) and per GST endpoint. M6, M19 and M20 are one bill's own figures
and don't change (one test shows a cancelled bill keeps them); M21 (search) is Task 7's."""

from decimal import Decimal
from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from billing.constants import BILL_CANCELLED
from billing.models import Product
from billing.services.gstr1 import document_series
from billing.tests.v3_helpers import ROLE_GROUPS, buyer, client_for, firm, line, person, sale

SEPT = "2026-09-10"


class CountedFiguresTest(TestCase):
    """One active ₹10,300 sale and one cancelled ₹20,600 sale: every figure is the active one's."""

    def setUp(self):
        self.client = client_for(person("kailash", *ROLE_GROUPS["owner"]))
        self.biz = firm()
        self.cust = buyer()
        Product.objects.create(name="Silver", hsn_code="711311", gst_tax_rate=Decimal("0.03"))
        self.live = sale(self.biz, self.cust, "1", SEPT)
        self.dead = sale(self.biz, self.cust, "2", SEPT, line(rate="20000"), status=BILL_CANCELLED)

    def get(self, name, **params):
        r = self.client.get(reverse(name), params)
        self.assertEqual(r.status_code, 200, getattr(r, "data", None))
        return r.data

    def test_m1_dashboard_stats(self):
        data = self.get("invoice-stats")
        self.assertEqual(data["totals"]["outward"], 10300.0)
        self.assertEqual(data["totals"]["count"], 1)
        self.assertEqual(data["totals"]["tax"], 300.0)
        self.assertEqual(data["monthly"][0]["outward_total"], 10300.0)
        self.assertEqual(data["top_customers"][0]["total"], 10300.0)
        self.assertEqual(data["top_products"][0]["total"], 10300.0)
        self.assertEqual(data["tax_distribution"]["cgst"], 150.0)
        # the recent list keeps the cancelled bill, marked
        self.assertEqual({r["invoice_number"]: r["status"] for r in data["recent_invoices"]},
                         {"1": "active", "2": BILL_CANCELLED})

    def test_m2_totals(self):
        self.assertEqual(self.get("invoice-totals")["outward_total"], Decimal("10300"))

    def test_m3_monthly_totals(self):
        self.assertEqual(self.get("invoice-monthly-totals")[0]["outward_total"], Decimal("10300"))

    def test_m4_distribution(self):
        data = self.get("invoice-distribution")
        self.assertEqual((data["outward_total"], data["outward_count"]), (Decimal("10300"), 1))

    def test_m5_data_quality(self):
        sale(self.biz, self.cust, "3", SEPT, status=BILL_CANCELLED).lineitem_set.all().delete()
        sale(self.biz, self.cust, "4", SEPT, line(hsn=""), status=BILL_CANCELLED)
        # A cancelled bill still holds its number, so it stays in a duplicate group. Two sales can't
        # share a number (uniq_outward_number_per_business_fy), so the pair is two purchases.
        sale(self.biz, self.cust, "SJ-7", SEPT, kind="inward")
        sale(self.biz, self.cust, "SJ-7", SEPT, kind="inward", status=BILL_CANCELLED)
        data = self.get("invoice-data-quality")
        self.assertEqual(data["invoices_no_line_items"], 0)
        self.assertEqual(data["line_items_missing_hsn"], 0)
        self.assertEqual(data["duplicate_invoice_groups"], 1)

    def test_m6_m19_m20_a_cancelled_bills_own_figures_stay(self):
        row = next(r for r in self.get("invoice-list")["results"] if r["id"] == self.dead.pk)
        self.assertEqual((row["total_tax"], row["line_item_count"]), ("600.00", 1))
        for name, key, value in (("invoice-summary", "total_tax", "600.00"), ("invoice-print", "total_amount", 20600),
                                 ("invoice-eway-bill", "total_amount", 20600.0)):
            r = self.client.get(reverse(name, args=[self.dead.pk]))
            self.assertEqual((r.status_code, r.data[key]), (200, value), name)

    def test_m7_firms_list(self):
        row = self.get("business-list")["results"][0]
        self.assertEqual((row["total_revenue"], row["invoice_count"]), (10300.0, 1))

    def test_m8_firms_performance(self):
        self.assertEqual(self.get("business-performance")[0]["outward_total"], Decimal("10300"))

    def test_m9_customers_list(self):
        row = self.get("customer-list")["results"][0]
        self.assertEqual((row["total_revenue"], row["invoice_count"]), (10300.0, 1))

    def test_m10_top_customers(self):
        self.assertEqual(self.get("customer-top")[0]["total_amount"], Decimal("10300"))

    def test_m11_products_list(self):
        self.assertEqual(self.get("product-list")["results"][0]["total_revenue"], 10300.0)

    def test_m12_top_products(self):
        self.assertEqual(self.get("product-top")[0]["total_amount"], Decimal("10300"))

    def test_m13_hsn_usage(self):
        product = Product.objects.get(name="Silver")
        r = self.client.get(reverse("product-hsn-usage", args=[product.pk]))
        self.assertEqual(r.data["variants"][0]["lines"], 1)

    def test_m14_gst_summary(self):
        data = self.get("invoice-gst-summary", business_id=self.biz.pk)
        self.assertEqual(data["rate_slabs"]["outward"][0]["taxable"], 10000.0)
        self.assertEqual(data["hsn_summary"][0]["taxable"], 10000.0)
        self.assertEqual(data["gstr3b"]["output_tax"]["total"], 300.0)

    def test_m15_gstr_export(self):
        data = self.get("invoice-gstr-export", business_id=self.biz.pk)
        self.assertEqual([row["txval"] for row in data["gstr1"]["b2cs"]], [10000.0])
        self.assertEqual(data["gstr3b"]["sup_details"]["osup_det"]["txval"], 10000.0)

    def test_m16_gstr1_portal_json(self):
        data = self.get("invoice-gstr1-portal-json", business_id=self.biz.pk, month=9, year=2026)
        self.assertEqual(data["meta"]["invoice_counts"]["b2cs"], 1)
        self.assertEqual(data["meta"]["taxable_total"], 10000.0)

    def test_m17_ca_excel(self):
        r = self.client.post(reverse("generate-report"), {"start_date": "2026-09-01", "end_date": "2026-09-30",
                                                          "invoice_type": "outward"}, format="json")
        sheet = load_workbook(BytesIO(r.content)).worksheets[0]
        grand = next(row for row in sheet.iter_rows(values_only=True) if "Grand Total" in row)
        self.assertEqual(Decimal(str(grand[10])), Decimal("10000"))

    def test_m18_reconciliation(self):
        data = self.get("reconciliation", fy="2026-27")
        self.assertEqual(data["total"]["invoice_count"], 1)
        self.assertEqual(data["total"]["gstr3b"]["taxable"], "10000.00")


class TableThirteenTest(TestCase):
    """A cancelled number stays in the month's series as cancelled (design decision 5)."""

    def setUp(self):
        self.client = client_for(person("kailash", *ROLE_GROUPS["owner"]))
        self.biz = firm()
        self.cust = buyer()
        self.bills = [sale(self.biz, self.cust, str(n), SEPT) for n in (1, 2, 3)]

    def docs(self):
        r = self.client.get(reverse("invoice-gstr1-portal-json"), {"business_id": self.biz.pk, "month": 9, "year": 2026})
        return r.data["file"]["doc_issue"]["doc_det"][0]["docs"]

    def test_a_cancelled_last_number_keeps_its_place(self):
        self.bills[2].status = BILL_CANCELLED
        self.bills[2].save()
        self.assertEqual(self.docs(), [{"num": 1, "from": "1", "to": "3", "totnum": 3, "cancel": 1, "net_issue": 2}])

    def test_a_cancelled_first_number_keeps_its_place(self):
        self.bills[0].status = BILL_CANCELLED
        self.bills[0].save()
        self.assertEqual(self.docs(), [{"num": 1, "from": "1", "to": "3", "totnum": 3, "cancel": 1, "net_issue": 2}])

    def test_document_series_without_cancelled_numbers_is_unchanged(self):
        self.assertEqual(document_series(["1", "3"]),
                         [{"num": 1, "from": "1", "to": "3", "totnum": 3, "cancel": 1, "net_issue": 2}])
        self.assertEqual(document_series(["A"], cancelled=["B"]),
                         [{"num": 1, "from": "A", "to": "A", "totnum": 1, "cancel": 0, "net_issue": 1},
                          {"num": 2, "from": "B", "to": "B", "totnum": 1, "cancel": 1, "net_issue": 0}])
