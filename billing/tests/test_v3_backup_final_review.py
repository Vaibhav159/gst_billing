"""Final part-4 regressions across backup, restore and v2 rollback."""
import json
from datetime import date
from io import BytesIO, StringIO

import pytest
from django.core.management import call_command
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.exceptions import ValidationError
from rest_framework_simplejwt.exceptions import TokenError
from rest_framework_simplejwt.token_blacklist.models import BlacklistedToken, OutstandingToken
from rest_framework_simplejwt.tokens import RefreshToken

from billing.models import Customer, Invoice, ITCReclaimLedger
from billing.services import backup_restore, backups
from billing.services.record_exports import response_for
from billing.tests.test_v3_backup_restore import apply, payload, purchase_payload, upload
from billing.tests.v3_helpers import buyer, client_for, firm, person, sale

pytestmark = pytest.mark.django_db


def current_payload():
    return {**payload("5.0"), "app": "gst-billing"}


def test_cancelled_sales_stay_out_of_v2_invoices_and_restore_in_v3():
    owner = person("backup-owner", "admin")
    business, customer = firm(), buyer()
    active = sale(business, customer, "1", "2026-03-01")
    cancelled = sale(business, customer, "2", "2026-03-01", status="cancelled", cancel_reason="Repeated")
    sale(business, customer, "P1", "2026-03-01", kind="inward")
    data = json.loads(b"".join(backups.full_stream(owner, timezone.now(), "backup.json")))
    assert [row["invoice_number"] for row in data["invoices"]] == [active.invoice_number, "P1"]
    assert [row["invoice_number"] for row in data["v3"]["cancelled"]] == [cancelled.invoice_number]
    assert data["counts"]["cancelled"] == 1
    Invoice.objects.all().delete()
    data["invoices"] = [row for row in data["invoices"] if row["type_of_invoice"] == "outward"]
    result = apply(data, owner)
    assert result["plan"]["bills"]["add_count"] == 2
    assert Invoice.objects.get(invoice_number="2").status == "cancelled"
    assert Invoice.objects.get(invoice_number="2").cancel_reason == "Repeated"


@pytest.mark.parametrize("cancelled", [{}, [None], ["bad"]])
def test_cancelled_block_is_validated(cancelled):
    data = current_payload()
    data["v3"] = {"cancelled": cancelled}
    with pytest.raises(ValidationError):
        backup_restore.preview(upload(data))


def test_cancelled_block_participates_in_year_selection():
    data = current_payload()
    row = {**data["invoices"][0], "status": "cancelled"}
    data["invoices"] = []
    data["v3"] = {"cancelled": [{**row, "invoice_number": str(n + 1)} for n in range(3001)]}
    answer = backup_restore.preview(upload(data))
    assert answer["years"][0]["bills"] == 3001
    assert answer["years"][0]["fy"] == "2025-26"


@pytest.mark.parametrize("purchase", [False, True])
@pytest.mark.parametrize("by_id", [False, True])
def test_same_name_different_registered_party_is_refused(purchase, by_id):
    owner = person("restore-owner", "admin")
    existing = buyer("New Buyer", gst_number="08VRHFA2322D1ZI")
    data = purchase_payload() if purchase else current_payload()
    data["customers"][0]["gst_number"] = "08WWUFP9975S1ZN"
    row = data["invoices"][0]
    row.update(customer_name="New Buyer", customerGST="08WWUFP9975S1ZN")
    if not by_id:
        data["customers"] = []
        row["customer"] = None
    result = apply(data, owner)
    kind = "purchases" if purchase else "bills"
    assert result["plan"][kind]["fail"][0]["why"] == "same name as New Buyer, but a different GSTIN"
    assert not Invoice.objects.exists()
    assert Customer.objects.get(pk=existing.pk).gst_number == "08VRHFA2322D1ZI"


@pytest.mark.parametrize("existing_gstin,file_gstin", [("", "08WWUFP9975S1ZN"), ("08VRHFA2322D1ZI", "")])
def test_name_fallback_is_kept_when_one_party_is_unregistered(existing_gstin, file_gstin):
    existing = buyer("New Buyer", gst_number=existing_gstin)
    data = current_payload()
    data["customers"][0]["gst_number"] = file_gstin
    owner = person("restore-owner", "admin")
    result = apply(data, owner)
    assert result["plan"]["bills"]["add_count"] == 1
    assert Invoice.objects.get().customer_id == existing.pk


def test_invalid_firm_reason_is_kept_on_dependent_bills():
    data = payload()
    data["businesses"][0]["gst_number"] = "BAD"
    plan = backup_restore.preview(upload(data))["plan"]
    assert plan["bills"]["fail"][0]["why"] == "Its GSTIN BAD doesn't check out"


@pytest.mark.parametrize("day", ["1990-01-01", "2090-01-01"])
def test_restored_sales_obey_sales_date_bounds(day):
    data = payload()
    data["invoices"][0]["invoice_date"] = day
    plan = backup_restore.preview(upload(data))["plan"]
    assert plan["bills"]["add_count"] == 0
    assert plan["bills"]["fail_count"] == 1


@pytest.mark.parametrize("items", [[], [{"product_name": "Gold", "quantity": "1", "rate": "100", "gst_tax_rate": "0.03", "hsn_code": ""}]])
def test_restored_sales_need_items_and_hsn(items):
    data = payload()
    data["invoices"][0]["line_items"] = items
    plan = backup_restore.preview(upload(data))["plan"]
    assert plan["bills"]["add_count"] == 0
    assert plan["bills"]["fail_count"] == 1


@pytest.mark.parametrize("kind", ["product", "sale"])
@pytest.mark.parametrize("rate", ["0.07", "25", "100"])
def test_restore_rejects_off_slab_tax_rates(kind, rate):
    data = payload()
    row = data["products"][0] if kind == "product" else data["invoices"][0]["line_items"][0]
    row["gst_tax_rate"] = rate
    with pytest.raises(ValidationError):
        backup_restore.preview(upload(data))


def test_restore_keeps_supplier_mark_without_a_purchase():
    data = payload()
    data["invoices"] = []
    data["customers"][0]["is_supplier"] = True
    apply(data, person("supplier-owner", "admin"))
    assert Customer.objects.get().is_supplier


def test_preledger_purchase_discards_old_credit_month():
    data = purchase_payload()
    data["businesses"][0]["gst_number"] = "08VRHFA2322D1ZI"
    data["customers"][0]["gst_number"] = "08WWUFP9975S1ZN"
    data["invoices"][0].update(invoice_date="2025-03-31", itc_month="2026-04-01")
    result = apply(data, person("credit-owner", "admin"))
    assert result["plan"]["purchases"]["add_count"] == 1
    assert Invoice.objects.get().itc_month is None


@pytest.mark.parametrize("purchase", [False, True])
def test_preview_flags_recalculated_tax_heads_and_totals(purchase):
    data = purchase_payload() if purchase else payload()
    data["businesses"][0]["gst_number"] = "08VRHFA2322D1ZI"
    data["customers"][0]["gst_number"] = "08WWUFP9975S1ZN"
    data["invoices"][0]["line_items"][0].update(cgst="0", sgst="0", igst="0.25")
    data["invoices"][0]["total_amount"] = "999"
    plan = backup_restore.preview(upload(data))["plan"]
    detail = plan["purchases" if purchase else "bills"]["add"][0]["detail"]
    assert "tax heads" in detail.lower()
    assert "total" in detail.lower()


@pytest.mark.parametrize("name", ["Bad\x0bName", "\x01=1+1"])
def test_xlsx_strips_illegal_controls_and_keeps_formula_protection(name):
    owner = person("export-owner", "admin")
    sale(firm(), buyer(name), "1", date(2026, 3, 1))
    response = response_for(owner, "sales", {"format": "xlsx", "fy": "all"})
    try:
        workbook = load_workbook(BytesIO(b"".join(response.streaming_content)))
        cell = workbook.active.cell(2, 4)
        assert cell.value == ("BadName" if name.startswith("Bad") else "'=1+1")
        assert cell.data_type == "s"
    finally:
        response.close()


def test_rollback_blacklists_every_token_only_on_forward_apply():
    first, second = person("first-owner", "admin"), person("second-owner", "admin")
    refreshes = [RefreshToken.for_user(first), RefreshToken.for_user(second)]
    assert OutstandingToken.objects.count() == 2
    call_command("prepare_v2_rollback", stdout=StringIO())
    assert not BlacklistedToken.objects.exists()
    call_command("prepare_v2_rollback", "--reverse", "--apply", stdout=StringIO())
    assert not BlacklistedToken.objects.exists()
    for _ in range(2):
        call_command("prepare_v2_rollback", "--apply", stdout=StringIO())
        assert BlacklistedToken.objects.count() == 2
    for refresh in refreshes:
        with pytest.raises(TokenError, match="blacklisted"):
            RefreshToken(str(refresh))


def test_v2_credit_get_returns_defaults_without_creating_history():
    business = firm()
    client = client_for(person("ledger-owner", "admin"))
    response = client.get(f"/api/itc-ledger/{business.pk}/")
    assert response.status_code == 200
    assert response.data["opening_cgst"] == "0.00"
    assert not ITCReclaimLedger.objects.exists()
    response = client.patch(f"/api/itc-ledger/{business.pk}/", {"opening_cgst": "2.00"}, format="json")
    assert response.status_code == 200
    assert ITCReclaimLedger.objects.get().opening_cgst == 2
    assert client.get(f"/api/itc-ledger/{business.pk}/").data["opening_cgst"] == "2.00"


def test_direct_bill_matches_gstin_before_a_mapped_name():
    buyer("New Buyer")
    registered = buyer("Registered buyer", gst_number="08WWUFP9975S1ZN")
    data = current_payload()
    data["invoices"][0].update(customer=None, customer_name="New Buyer", customerGST=registered.gst_number)
    result = apply(data, person("gstin-first-owner", "admin"))
    assert result["plan"]["bills"]["add_count"] == 1
    assert Invoice.objects.get().customer_id == registered.pk


@pytest.mark.parametrize("problem", ["empty", "hsn", "slab", "date"])
def test_backup_refuses_invalid_authoritative_bin_contents(problem):
    from billing.services.bin import bin_bill

    owner = person("bin-validation-owner", "admin")
    business, customer = firm(), buyer()
    day = "1999-12-01" if problem == "date" else "2026-03-01"
    binned, _ = bin_bill(sale(business, customer, "1", day), owner)
    if problem == "empty":
        binned.data["lines"] = []
    elif problem == "hsn":
        binned.data["lines"][0]["hsn_code"] = ""
    elif problem == "slab":
        binned.data["lines"][0]["gst_tax_rate"] = "0.07"
    binned.save(update_fields=["data"])
    data = payload()
    data["businesses"][0]["name"] = business.name
    data["customers"][0]["name"] = customer.name
    data["invoices"][0].update(invoice_number="1", invoice_date="2000-01-01" if problem == "date" else "2026-03-01")
    result = apply(data, owner)
    assert result["plan"]["bills"]["fail_count"] == 1
    assert not Invoice.objects.exists()


def test_backup_validates_saved_bin_items_instead_of_obsolete_file_items():
    from billing.services.bin import bin_bill

    owner = person("bin-authority-owner", "admin")
    business, customer = firm(), buyer()
    binned, _ = bin_bill(sale(business, customer, "1", "2026-03-01"), owner)
    data = payload()
    data["businesses"][0]["name"] = business.name
    data["customers"][0]["name"] = customer.name
    data["invoices"][0].update(invoice_number="1", line_items=[])
    result = apply(data, owner)
    assert result["plan"]["bills"]["add_count"] == 1
    restored = Invoice.objects.get(pk=binned.original_id)
    assert restored.lineitem_set.count() == 1
    assert restored.total_amount == 10300


@pytest.mark.parametrize("existing_firm", [False, True])
@pytest.mark.parametrize("source", ["missing", "null", "historical"])
def test_backup_sales_keep_firm_snapshots_after_switch(existing_firm, source):
    from unittest.mock import patch

    from django.db import connection
    from django.test.utils import CaptureQueriesContext

    from billing.services.sales import firm_snapshot

    owner = person("snapshot-restore-owner", "admin")
    business = firm("NEW GOLD", gstin="", bank_name="Current bank") if existing_firm else None
    call_command("freeze_firm_snapshots", "--apply", stdout=StringIO())
    data = current_payload()
    if source != "missing":
        data["invoices"][0]["firm_snapshot"] = None if source == "null" else {"name": "Earlier firm", "bank_name": "Earlier bank"}
    data["invoices"].append({**data["invoices"][0], "invoice_number": "2"})
    with (patch.object(backup_restore, "firm_snapshot", wraps=firm_snapshot, create=True) as frozen,
          CaptureQueriesContext(connection) as queries):
        apply(data, owner)
    snapshots = list(Invoice.objects.order_by("pk").values_list("firm_snapshot", flat=True))
    if source == "historical":
        assert snapshots == [{"name": "Earlier firm", "bank_name": "Earlier bank", "signature_image": ""}] * 2
        assert frozen.call_count == 0
    else:
        business = business or Invoice.objects.first().business
        assert snapshots == [firm_snapshot(business)] * 2
        assert frozen.call_count == 1
    # All print fields arrive together; no per-attribute deferred Business fetch.
    assert not any('FROM "billing_business"' in query["sql"] and '"billing_business"."address"' in query["sql"]
                   and '"billing_business"."name"' not in query["sql"] for query in queries)
