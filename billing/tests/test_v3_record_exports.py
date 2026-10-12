import csv
from io import BytesIO, StringIO
from types import SimpleNamespace
from unittest.mock import Mock, patch

from django.test import TransactionTestCase
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from billing.models import AuditLog, LineItem, Product
from billing.tests.v3_helpers import buyer, firm, line, person, sale


class RecordExportTests(TransactionTestCase):
    """Response.close sends request_finished, so these stream tests need autocommit."""

    def setUp(self):
        self.viewer = person("viewer", "viewer")
        self.biz = firm(invoice_prefix="KGH")
        self.customer = buyer("=Danger", mobile_number="+9999999999")
        self.customer.businesses.add(self.biz)
        self.bill = sale(self.biz, self.customer, "KGH/2026-27/1", "2026-10-08", line(rate="100"))
        sale(self.biz, self.customer, "KGH/2026-27/2", "2026-10-09", line(rate="200"), status="cancelled")
        sale(self.biz, self.customer, "OLD", "2026-03-31", line(rate="500"))
        sale(self.biz, self.customer, "P1", "2026-10-10", line(rate="300"), kind="inward")
        Product.objects.create(name="@Silver", hsn_code="711311", gst_tax_rate="0.0025")

    def call(self, kind, **params):
        from billing.api.backups import RecordExportView
        request = APIRequestFactory().get("/api/exports/" + kind + "/", params)
        force_authenticate(request, user=self.viewer)
        return RecordExportView.as_view()(request, kind=kind)

    def rows(self, response):
        self.assertEqual(response.status_code, 200, getattr(response, "data", None))
        return list(csv.reader(StringIO(b"".join(response.streaming_content).decode("utf-8-sig"))))

    def test_sales_columns_formula_protection_and_cancelled_excluded_from_totals(self):
        response = self.call("sales", format="csv", fy="2026-27")
        self.assertEqual(response["Content-Disposition"], 'attachment; filename="Sales_FY2026-27_All-firms.csv"')
        rows = self.rows(response)
        self.assertEqual(rows[0], ["Date", "Bill no.", "Firm", "Customer", "GSTIN", "State", "Payment",
                                  "Taxable", "CGST", "SGST", "IGST", "Total", "Status", "From paper book"])
        self.assertEqual(rows[1][3], "'=Danger")
        self.assertEqual(rows[2][12], "Cancelled")
        self.assertEqual(rows[-1][7:12], ["100.00", "1.50", "1.50", "0.00", "103.00"])
        self.assertEqual(len(rows), 4)
        self.assertIn("Sales bills · FY 2026-27 · All firms · 2 bills", AuditLog.objects.get().details)

    def test_dates_and_firm_control_rows_and_filename(self):
        response = self.call("sales", format="csv", start_date="2026-03-31", end_date="2026-03-31", business_id=self.biz.pk)
        self.assertIn("Sales_2026-03-31_to_2026-03-31_Kiran.csv", response["Content-Disposition"])
        self.assertEqual(self.rows(response)[1][1], "OLD")

    def test_all_years_and_xlsx_stream_has_no_formulas(self):
        response = self.call("sales", format="xlsx", fy="all")
        self.assertTrue(response.streaming)
        book = load_workbook(BytesIO(b"".join(response.streaming_content)))
        rows = list(book.active.values)
        self.assertEqual(len(rows), 5)
        self.assertEqual(rows[-1][7:12], (600, 9, 9, 0, 618))
        self.assertEqual(rows[1][3], "'=Danger")
        self.assertFalse(any(cell.data_type == "f" for row in book.active for cell in row))
        response.close()

    def test_formula_prefixes_are_protected_in_csv_and_xlsx(self):
        unsafe = ("\t=1+1", "\r=1+1", " =1+1", " \t\r+1", "  -1", "  @SUM(1,1)",
                  "\tplain", "\rplain")
        for value in (*unsafe, "  Jewellery shop"):
            self.customer.name = value
            self.customer.save(update_fields=["name"])
            expected = "'" + value if value in unsafe else value
            for format_name in ("csv", "xlsx"):
                with self.subTest(value=value, format=format_name):
                    response = self.call("sales", format=format_name, fy="2026-27")
                    try:
                        if format_name == "csv":
                            rows = self.rows(response)
                            self.assertEqual(rows[1][3], expected)
                            self.assertEqual(rows[1][11], "103.00")
                        else:
                            book = load_workbook(BytesIO(b"".join(response.streaming_content)))
                            # XML readers normalize carriage returns to line feeds.
                            self.assertEqual(book.active.cell(2, 4).value, expected.replace("\r", "\n"))
                            self.assertEqual(book.active.cell(2, 4).data_type, "s")
                            self.assertEqual(book.active.cell(2, 12).value, 103)
                            self.assertEqual(book.active.cell(2, 12).data_type, "n")
                    finally:
                        response.close()

    def test_purchase_columns_and_part2_credit(self):
        rows = self.rows(self.call("purchases", format="csv", fy="2026-27"))
        self.assertEqual(rows[0], ["Date", "Supplier bill no.", "Firm", "Supplier", "GSTIN", "Taxable",
                                  "CGST", "SGST", "IGST", "Total", "Credit month", "Credit allowed"])
        self.assertEqual(rows[1][1], "P1")
        # 2A's credit month: a bill dated October 2026 with no chosen month credits that month.
        self.assertEqual(rows[1][-2:], ["2026-10", "Yes"])
        self.assertEqual(rows[-1][5:10], ["300.00", "4.50", "4.50", "0.00", "309.00"])

    def test_customer_columns_and_formula_phone(self):
        rows = self.rows(self.call("customers"))
        self.assertEqual(rows[0], ["Name", "Type", "Phone", "Email", "GSTIN", "PAN", "State", "City", "Address", "Firms"])
        self.assertEqual(rows[1][0], "'=Danger")
        self.assertEqual(rows[1][2], "'+9999999999")
        self.assertEqual(rows[1][-1], "KIRAN GOLD HOUSE")

    def test_product_columns_and_fractional_rate(self):
        kind_reader = Mock(return_value="silver")
        with patch.dict("sys.modules", {"billing.services.catalog": SimpleNamespace(effective_kind=kind_reader)}):
            rows = self.rows(self.call("products"))
        kind_reader.assert_called_once()
        self.assertIsNone(kind_reader.call_args.args[0].kind)
        self.assertEqual(rows[0], ["Name", "HSN", "GST rate", "Unit", "Kind", "Purity", "Usual rate", "Bought only", "Description"])
        self.assertEqual(rows[1][:5], ["'@Silver", "711311", "0.25", "gms", "silver"])

    def test_firm_columns_and_plain_bank_account(self):
        rows = self.rows(self.call("firms"))
        self.assertEqual(rows[0], ["Name", "Short name", "GSTIN", "PAN", "State", "City", "PIN", "Address",
                                  "Phone", "Email", "Prefix", "Bank", "Account number", "IFSC", "Branch"])
        self.assertEqual(rows[1][0], "KIRAN GOLD HOUSE")

    def test_bad_inputs_are_400_in_exact_words(self):
        cases = [("other", {}, "kind", "Choose sales, purchases, customers, products or firms."),
                 ("sales", {"format": "pdf"}, "format", "Choose xlsx or csv."),
                 ("products", {"format": "xlsx"}, "format", "Choose xlsx or csv."),
                 ("sales", {"fy": "2026-99"}, "fy", "Give the financial year like 2026-27."),
                 ("sales", {"start_date": "wrong"}, "start_date", "Type the date like 2026-09-30."),
                 ("sales", {"start_date": "2026-10-09", "end_date": "2026-10-08"}, "end_date", "The start date is after the end date."),
                 ("sales", {"business_id": "999999999999999999999"}, "business_id", "Pick a firm from the list.")]
        for kind, params, field, words in cases:
            with self.subTest(kind=kind, params=params):
                response = self.call(kind, **params)
                self.assertEqual(response.status_code, 400)
                self.assertEqual(response.data, {field: [words]})

    def test_export_success_is_not_audited_when_csv_stream_abandoned(self):
        response = self.call("sales", format="csv", fy="all")
        next(iter(response.streaming_content))
        response.close()
        self.assertFalse(AuditLog.objects.exists())


    def test_legacy_heads_round_each_line_before_bill_and_export_totals(self):
        tiny = sale(self.biz, self.customer, "TINY", "2026-11-01", line(rate="1"), line(rate="1"))
        LineItem.objects.filter(invoice=tiny).update(cgst="0.025", sgst="0.025", igst=0)
        for kind in ("sales", "purchases"):
            from billing.models import Invoice
            Invoice.objects.filter(pk=tiny.pk).update(type_of_invoice="outward" if kind == "sales" else "inward")
            rows = self.rows(self.call(kind, format="csv", start_date="2026-11-01", end_date="2026-11-01"))
            offset = 7 if kind == "sales" else 5
            self.assertEqual(rows[1][offset:offset + 5], ["2.00", "0.06", "0.06", "0.00", "2.12"])
            self.assertEqual(rows[-1][offset:offset + 5], ["2.00", "0.06", "0.06", "0.00", "2.12"])


    def test_published_audit_fixtures_match_every_download_producer(self):
        import json
        from pathlib import Path

        from billing.api.backups import BackupFullView

        request = APIRequestFactory().get("/api/backup/full/")
        force_authenticate(request, user=person("backup-fixture-owner", "admin"))
        b"".join(BackupFullView.as_view()(request).streaming_content)
        with patch.dict("sys.modules", {"billing.services.catalog": SimpleNamespace(effective_kind=lambda product: "silver")}):
            for kind in ("sales", "purchases", "customers", "products", "firms"):
                b"".join(self.call(kind, format="csv", fy="all").streaming_content)
        expected = json.loads((Path(__file__).parent / "fixtures/v3_backup_audit.json").read_text())
        keys = ("entity", "action", "entity_id", "entity_name")
        actual = list(AuditLog.objects.order_by("pk").values(*keys))
        self.assertEqual(actual, [{key: row[key] for key in keys} for row in expected])


    def test_purchase_credit_month_uses_shared_reader_for_null_storage(self):
        from datetime import date

        from billing.services.record_exports import bill_rows

        for day, effective in ((date(2024, 3, 31), None), (date(2026, 10, 8), date(2026, 10, 1))):
            projected = {"id": self.bill.pk, "invoice_date": day, "invoice_number": "P1",
                         "business__name": self.biz.name, "customer__name": self.customer.name,
                         "customer__gst_number": "", "status": "active", "itc_month": None,
                         "itc_eligible": True}
            # 2A's real credit_month_of: before April 2025 a blank month is outside the ledger.
            with patch("billing.services.record_exports.batches", return_value=iter([[projected]])):
                rows = list(bill_rows("purchases", None, None, None, [0]))
            self.assertEqual(rows[0][-2], effective.strftime("%Y-%m") if effective else "")
